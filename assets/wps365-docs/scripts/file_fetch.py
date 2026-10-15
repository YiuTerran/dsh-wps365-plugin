#!/usr/bin/env python3
"""读取 WPS 365 云文件：下载原件 + 抽出可读文本。

统一套路（.docx / .xlsx / .pptx 都适用）：
    file_fetch.py 下载原件 → 本地用 Office 技能编辑 → file_write.py 传回

本脚本负责第一步，并顺手把内容抽成 Markdown 供 agent 直接阅读：

  .docx  GET .../content?format=markdown&include_elements=all → data.markdown
  .xlsx  本地 openpyxl 解析（**能拿到公式原文**、精确值、合并区、多表）
         `--via-api` 改用 sheets 接口，不下载，但拿不到公式原文且显示值会丢精度
  其他   只下载原件

关键点：
  * .docx 的正文接口**必须带 `include_elements=all`**，否则表格静默消失
  * .xlsx 的 sheets 接口**读不到公式原文**（`=B2*C2` 只回 `16.5`），且 `cell_text`
    是显示值会丢精度（`1234567890123` → `1.23457E+12`）。所以默认走本地解析。
  * 下载必须带 delegated token，否则 `{"result":"userNotLogin"}`

产物：
  <out>/doc.md          可读文本（.docx 正文 / .xlsx 逐表表格 + 公式清单）
  <out>/meta.json       元数据（name/version/size/link_url）
  <out>/original.<ext>  原件字节（--download；.xlsx 默认就要下载）
  <out>/<sheet>.csv     每个工作表一份（--csv）
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import urllib.request

from _wps_cli import resolve as _resolve_cli
from _wps_cli import run_cli as _run_cli
from _wps_cli import use_utf8_stdio as _use_utf8_stdio

DEFAULT_CLI = _resolve_cli()
_LABEL = "file_fetch"
_use_utf8_stdio()
UA = {"User-Agent": "Mozilla/5.0"}
SHEET_ROWS_PER_CALL = 100  # sheets 接口按行分块，避免单次响应过大


# ---------------------------------------------------------------- CLI 封装

def run(cmd: list[str], timeout: int, what: str) -> str:
    return _run_cli(cmd, timeout, _LABEL, what)


def api(cli: str, method: str, path: str, timeout: int,
        query: dict | None = None, data: dict | None = None) -> dict:
    """调用 `wps365-cli api`，返回 data 段。

    超时必须同时传给 CLI 本身 —— 它有自己的 HTTP 超时（默认较短）。
    """
    cmd = [cli, "api", method, path, "-o", "json", "--timeout", f"{timeout}s",
           "--token-type", "delegated"]
    for key, value in (query or {}).items():
        cmd += ["--query", f"{key}={value}"]
    if data is not None:
        cmd += ["--data", json.dumps(data, ensure_ascii=False)]
    try:
        payload = json.loads(run(cmd, timeout, f"{method.upper()} {path}"))
    except json.JSONDecodeError as exc:
        sys.exit(f"[file_fetch] {method.upper()} {path} 返回不是 JSON：{exc}")
    if payload.get("code") not in (0, None):
        sys.exit(f"[file_fetch] {method.upper()} {path} 报错：{payload.get('msg') or payload}")
    return payload.get("data") or {}


def json_cmd(cli: str, args: list[str], timeout: int, what: str) -> dict:
    payload = json.loads(run([cli, *args, "-o", "json", "--timeout", f"{timeout}s"], timeout, what))
    if payload.get("code") not in (0, None):
        sys.exit(f"[file_fetch] {what} 报错：{payload.get('msg') or payload}")
    return payload.get("data") or {}


def auth_token(cli: str, timeout: int) -> str:
    return run([cli, "auth", "token"], timeout, "取 access token").strip()


# ---------------------------------------------------------------- 定位

def link_to_ids(cli: str, link: str, timeout: int) -> tuple[str, str]:
    """把 kdocs/plus 分享链接解析成 (drive_id, file_id)。"""
    link_id = link.rstrip("/").rsplit("/", 1)[-1].split("?")[0]
    data = api(cli, "get", f"/v7/links/{link_id}/meta", timeout)
    drive_id, file_id = data.get("drive_id"), data.get("file_id")
    if not drive_id or not file_id:
        sys.exit(f"[file_fetch] 链接解析结果缺少 drive_id/file_id：{data}")
    return drive_id, file_id


def drive_of_file(cli: str, file_id: str, timeout: int) -> str:
    """只给了 file_id 时，反查它所在的盘。"""
    data = api(cli, "get", f"/v7/files/{file_id}/meta", timeout)
    for key in ("drive_id", "drive"):
        value = data.get(key)
        if isinstance(value, dict):
            value = value.get("id")
        if value:
            return str(value)
    sys.exit(f"[file_fetch] 无法判断所在盘，请显式传 --drive-id：{data}")


def download_original(cli: str, drive_id: str, file_id: str, out_dir: str, name: str,
                      timeout: int) -> tuple[str, int, str, str]:
    """下载原件，返回 (路径, 字节数, 本地 sha1, 接口声明 sha1)。"""
    info = json_cmd(cli, ["drive", "file", "download", drive_id, file_id, "--with-hash"],
                    timeout, "取下载地址")
    url = info.get("url")
    if not url:
        sys.exit("[file_fetch] 下载响应里没有 url")
    declared = next((h.get("sum") for h in (info.get("hashes") or [])
                     if str(h.get("type")).lower() == "sha1"), "")

    ext = os.path.splitext(name or "")[1] or ".bin"
    path = os.path.join(out_dir, "original" + ext)
    req = urllib.request.Request(url, headers={**UA,
                                              "Authorization": f"Bearer {auth_token(cli, timeout)}"})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        blob = resp.read()
    with open(path, "wb") as fh:
        fh.write(blob)
    return path, len(blob), hashlib.sha1(blob).hexdigest(), declared


# ---------------------------------------------------------------- .docx

def fetch_docx_markdown(cli: str, drive_id: str, file_id: str, timeout: int) -> tuple[str, bool]:
    """取正文 Markdown 与 is_partly_exported 标记。"""
    data = api(cli, "get", f"/v7/drives/{drive_id}/files/{file_id}/content", timeout,
               query={"format": "markdown", "include_elements": "all"})
    return (data.get("markdown") or data.get("plain") or ""), bool(data.get("is_partly_exported"))


def table_count(md: str) -> int:
    """数 Markdown 表格块（以分隔行 `| --- |` 为标志，比数竖线可靠）。"""
    return len(re.findall(r"(?m)^\s*\|[\s:|-]+\|\s*$", md))


# ---------------------------------------------------------------- .xlsx

def _esc(value) -> str:
    return str(value).replace("|", "\\|").replace("\n", " ").strip()


def xlsx_read_local(path: str, out_dir: str, max_rows: int, max_cols: int,
                    want_csv: bool) -> tuple[str, int, int]:
    """本地 openpyxl 解析：能拿到公式原文、精确值、合并区。

    返回 (markdown, 公式数, 工作表数)。
    """
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("[file_fetch] 本地解析 .xlsx 需要 openpyxl。请用工作区自带的 Python 运行"
                 "（load_workspace_dependencies 给出的可执行文件），或加 --via-api 改走 sheets 接口。")

    # 扩展名会骗人：云盘上带 .xlsx 名字的老式 .xls（OLE2 复合文档）并不少见，
    # openpyxl 会抛一句"File is not a zip file"让人摸不着头脑。这里先按魔数判断，
    # 给一条能直接照做的提示，而不是让异常栈直接冒到模型面前。
    with open(path, "rb") as fh:
        magic = fh.read(8)
    if magic == bytes.fromhex("d0cf11e0a1b11ae1"):
        sys.exit(f"[file_fetch] {os.path.basename(path)} 的名字是 .xlsx，内容却是老式 .xls"
                 "（OLE2 复合文档），openpyxl 打不开。请让用户另存为 .xlsx 后重试；"
                 "只想拿单元格值可以加 --via-api 走 sheets 接口（代价是拿不到公式原文）。")

    wb = load_workbook(path, data_only=False)      # 公式原文
    wbv = load_workbook(path, data_only=True)      # 缓存计算值
    lines: list[str] = []
    formulas: list[str] = []
    truncated: list[str] = []

    for name in wb.sheetnames:
        ws, wsv = wb[name], wbv[name]
        total_r = ws.max_row or 0
        total_c = ws.max_column or 0
        empty = all(c.value is None for row in ws.iter_rows() for c in row)
        max_r = min(total_r, max_rows)
        max_c = min(total_c, max_cols)
        if total_r > max_r or total_c > max_c:
            truncated.append(f"{name}（{total_r}×{total_c} 截到 {max_r}×{max_c}）")

        lines.append(f"## 工作表：{name}（{total_r} 行 × {total_c} 列）")
        if empty:
            lines += ["（空表）", ""]
            continue
        if ws.merged_cells.ranges:
            lines.append("合并区：" + ", ".join(str(r) for r in ws.merged_cells.ranges))
        lines.append("")

        letters = [get_column_letter(c) for c in range(1, max_c + 1)]
        lines.append("| 行 | " + " | ".join(letters) + " |")
        lines.append("| --- | " + " | ".join(["---"] * max_c) + " |")
        for r in range(1, max_r + 1):
            cells = []
            for c in range(1, max_c + 1):
                cell = ws.cell(row=r, column=c)
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    cached = wsv.cell(row=r, column=c).value
                    formulas.append(f"{name}!{cell.coordinate} = {v}"
                                    + (f"  → {cached}" if cached is not None else "  → (无缓存值)"))
                    v = f"{v} → {cached}" if cached is not None else v
                cells.append(_esc("" if v is None else v))
            lines.append(f"| {r} | " + " | ".join(cells) + " |")
        lines.append("")

        if want_csv:
            import csv
            csv_path = os.path.join(out_dir, f"{name}.csv")
            with open(csv_path, "w", newline="", encoding="utf-8-sig") as fh:
                writer = csv.writer(fh)
                for row in wsv.iter_rows(max_row=max_r, max_col=max_c, values_only=True):
                    writer.writerow(["" if v is None else v for v in row])

    if formulas:
        lines.append("## 公式清单（API 读不到，只有本地解析能拿到）")
        lines += [f"- {f}" for f in formulas]
        lines.append("")
    if truncated:
        lines.append("⚠️ 已按上限截断：" + "；".join(truncated)
                     + "（用 --max-rows / --max-cols 调整）")
    return "\n".join(lines), len(formulas), len(wb.sheetnames)


def xlsx_read_api(cli: str, file_id: str, timeout: int, max_rows: int,
                  max_cols: int) -> tuple[str, int, int]:
    """sheets 接口读取：不下载即可取值。

    两个固有缺陷（所以只是备选）：
      * 拿不到公式原文，只给计算结果
      * `cell_text` 是显示值会丢精度，这里一律用 `original_cell_value`
    """
    sheets = api(cli, "get", f"/v7/sheets/{file_id}/worksheets", timeout).get("sheets") or []
    lines: list[str] = []
    for sheet in sheets:
        sid, name = sheet.get("sheet_id"), sheet.get("name")
        area = sheet.get("active_area") or {}
        r0, r1 = area.get("row_from", 0), area.get("row_to", 0)
        c0, c1 = area.get("col_from", 0), area.get("col_to", 0)
        r1 = min(r1, r0 + max_rows - 1)
        c1 = min(c1, c0 + max_cols - 1)

        grid: dict[tuple[int, int], str] = {}
        r = r0
        while r <= r1:
            chunk_to = min(r + SHEET_ROWS_PER_CALL - 1, r1)
            data = api(cli, "get", f"/v7/sheets/{file_id}/worksheets/{sid}/range_data", timeout,
                       query={"row_from": r, "row_to": chunk_to, "col_from": c0, "col_to": c1})
            for cell in data.get("range_data") or []:
                value = cell.get("original_cell_value")
                if value in (None, ""):
                    value = cell.get("cell_text")
                if value not in (None, ""):
                    grid[(cell["row_from"], cell["col_from"])] = _esc(value)
            r = chunk_to + 1

        lines.append(f"## 工作表：{name}（sheet_id={sid}，使用范围 行 {r0}-{r1}、列 {c0}-{c1}）")
        lines.append("")
        if not grid:
            lines += ["（无数据）", ""]
            continue
        lines.append("| 行 | " + " | ".join(str(c) for c in range(c0, c1 + 1)) + " |")
        lines.append("| --- | " + " | ".join(["---"] * (c1 - c0 + 1)) + " |")
        for row_idx in range(r0, r1 + 1):
            values = [grid.get((row_idx, c), "") for c in range(c0, c1 + 1)]
            if any(values):
                lines.append(f"| {row_idx} | " + " | ".join(values) + " |")
        lines.append("")
    lines.append("注：本结果经 sheets 接口读取，**不含公式原文**，数值取自 "
                 "`original_cell_value`（`cell_text` 是显示值会丢精度）。")
    return "\n".join(lines), 0, len(sheets)


# ---------------------------------------------------------------- 主流程

def main() -> None:
    ap = argparse.ArgumentParser(description="读取 WPS 365 云文件（下载原件 + 抽文本）")
    ap.add_argument("--file-id", help="文件 file_id")
    ap.add_argument("--drive-id", help="所在盘 drive_id（缺省时按 file_id 反查）")
    ap.add_argument("--link", help="kdocs/plus 分享链接（与 --file-id 二选一）")
    ap.add_argument("--out", required=True, help="输出目录")
    ap.add_argument("--download", action="store_true", help="下载原件字节（.xlsx 总是下载）")
    ap.add_argument("--via-api", action="store_true",
                    help=".xlsx 改用 sheets 接口读取（不下载，但拿不到公式原文）")
    ap.add_argument("--csv", action="store_true", help=".xlsx 额外为每个工作表导出 CSV")
    ap.add_argument("--max-rows", type=int, default=200, help=".xlsx 最多解析行数")
    ap.add_argument("--max-cols", type=int, default=50, help=".xlsx 最多解析列数")
    ap.add_argument("--print-text", action="store_true", help="把 doc.md 打到 stdout（长文件慎用）")
    ap.add_argument("--cli", default=DEFAULT_CLI, help=f"wps365-cli 路径（默认 {DEFAULT_CLI}）")
    ap.add_argument("--timeout", type=int, default=180, help="单次 HTTP 超时秒数")
    args = ap.parse_args()

    if not args.file_id and not args.link:
        ap.error("需要 --file-id 或 --link")
    if not os.path.exists(args.cli):
        sys.exit(f"[file_fetch] 找不到 CLI：{args.cli}")

    drive_id, file_id = (link_to_ids(args.cli, args.link, args.timeout) if args.link
                         else (args.drive_id or "", args.file_id))
    if not drive_id:
        drive_id = drive_of_file(args.cli, file_id, args.timeout)

    meta = json_cmd(args.cli, ["drive", "file", "get", drive_id, file_id],
                    args.timeout, "取文件信息")
    name = meta.get("name") or ""
    ext = os.path.splitext(name)[1].lower()
    out_dir = os.path.abspath(args.out)
    os.makedirs(out_dir, exist_ok=True)

    is_xlsx = ext in (".xlsx", ".xls")
    need_binary = args.download or (is_xlsx and not args.via_api)

    dl = None
    if need_binary:
        dl = download_original(args.cli, drive_id, file_id, out_dir, name, args.timeout)

    md, extra, partial = "", "", False
    if ext in (".docx", ".doc"):
        md, partial = fetch_docx_markdown(args.cli, drive_id, file_id, args.timeout)
        tables = table_count(md)
        extra = (f"正文 {len(md)} 字符，识别到 {tables} 个表格"
                 + ("" if tables else "（若原文档确实有表格，说明 include_elements 没生效）"))
    elif is_xlsx:
        if args.via_api:
            md, _, n_sheet = xlsx_read_api(args.cli, file_id, args.timeout,
                                           args.max_rows, args.max_cols)
            extra = f"{n_sheet} 个工作表（经 API 读取，无公式原文）"
        else:
            md, n_formula, n_sheet = xlsx_read_local(dl[0], out_dir, args.max_rows,
                                                     args.max_cols, args.csv)
            extra = f"{n_sheet} 个工作表，{n_formula} 个公式（本地解析，含公式原文）"
    else:
        extra = f"未做文本抽取（{ext or '无扩展名'}），仅下载原件" if dl else "未做文本抽取"
        md = f"# {name}\n\n（该格式未做文本抽取，请直接使用 original{ext}）\n"

    md_path = os.path.join(out_dir, "doc.md")
    with open(md_path, "w", encoding="utf-8") as fh:
        fh.write(md)

    meta_path = os.path.join(out_dir, "meta.json")
    with open(meta_path, "w", encoding="utf-8") as fh:
        json.dump({"drive_id": drive_id, "file_id": file_id, "name": name,
                   "version": meta.get("version"), "size": meta.get("size"),
                   "link_url": meta.get("link_url"), "parent_id": meta.get("parent_id"),
                   # 基线：改完传回时用 --expect-hash / --expect-version 比对，
                   # 若期间被别人改过就会放弃更新（见 file_write.py 的冲突检查）
                   "hash": (meta.get("hash") or {}).get("sum"),
                   "hash_type": (meta.get("hash") or {}).get("type"),
                   "mtime": meta.get("mtime"),
                   "modified_by": (meta.get("modified_by") or {}).get("name")},
                  fh, ensure_ascii=False, indent=1)

    print(f"文件：{name or '(无名)'}  version={meta.get('version')}  size={meta.get('size')}")
    h = (meta.get("hash") or {})
    if h.get("sum"):
        print(f"基线：hash={h['sum']}（{h.get('type')}）"
              f"  最后修改={((meta.get('modified_by') or {}).get('name')) or '未知'}")
        if ext == ".otl":
            print("      → .otl 请用**块级局部比对**：otl_write.py --snapshot-out … --scope-index …，"
                  "再 --expect-snapshot 校验（整文件哈希对它没有局部意义）")
        else:
            print(f"      → 改完传回时用 --expect-hash {h['sum']} "
                  f"--expect-version {meta.get('version')} 防并发覆盖")
    print(f"链接：{meta.get('link_url') or '(无)'}")
    print(f"内容：{extra}")
    if partial:
        print("⚠️  服务端标记 is_partly_exported=true —— 只导出了一部分，正文可能不完整")
    if dl:
        _, nbytes, sha1, declared = dl
        match = ("✅ 一致" if declared and sha1 == declared
                 else ("⚠️ 与接口声明不同" if declared else "（接口未给 sha1）"))
        print(f"原件：{nbytes} 字节，本地 sha1={sha1[:16]}… {match}")
    print("产物：")
    print(f"  {md_path}  ({os.path.getsize(md_path)} 字节 —— 用 read 分页读取)")
    print(f"  {meta_path}")
    if dl:
        print(f"  {dl[0]}")
    if args.print_text:
        print("\n" + "=" * 60 + "\n")
        print(md)


if __name__ == "__main__":
    main()
