#!/usr/bin/env python3
"""改云端 .xlsx 的一条龙：下载 → 本地 openpyxl 编辑 → LibreOffice 重算 → 原地传回。

为什么不是调 sheets 接口：`.xlsx` 的单元格写接口**写不了字面值**（body 里只有
`formula`，传 `cell_text` 一律 `400000003`），且读不到公式原文。本地编辑没有这些
限制，公式、格式、合并、多工作表全部保留。

四步：
  1. 下载原件（同 file_fetch.py）
  2. 应用改动：`--py 脚本.py` 和/或 `--set 表!单元格=值`
  3. **重算公式缓存值**：openpyxl 保存会把 `data_only=True` 的缓存值抹掉，
     必须用 LibreOffice 重算一遍（无公式时自动跳过）
  4. 传回（同 file_write.py）：带 file_id = 原地替换，file_id 与链接不变

改动怎么写：
  --set "数据!B2=99"                  数字自动识别
  --set "数据!A11=新增文本"            其余按字符串
  --set "数据!D11==B11*C11"            值以 = 开头即公式（注意是两个等号）
  --set "第二页!A1=true"               布尔自动识别

  --py edit.py                        脚本以 wb / sheet() 为上下文执行：
                                      wb 是 Workbook（data_only=False）
                                      sheet("数据") 取工作表，不带参数取活动表
                                      可给 changes 列表用于打印改动摘要
"""
from __future__ import annotations

import argparse
from decimal import Decimal, InvalidOperation
import glob
import os
import shutil
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import file_fetch  # noqa: E402  下载 + 元数据
import file_write  # noqa: E402  三步上传 + 备份
from _wps_cli import CLI_ENCODING  # noqa: E402  管道一律按 UTF-8 解码
from _wps_cli import use_utf8_stdio as _use_utf8_stdio  # noqa: E402

_use_utf8_stdio()

#: DSH application bundles, by platform. The kit ships *inside* the application
#: on every platform, so the anchor is the app resource directory, not a
#: fixed POSIX path; the environment overrides below cover portable installs.
APP_RESOURCE_CANDIDATES = [
    "/Applications/DeepSeek Harness.app/Contents/Resources",
]
if os.name == "nt":
    for _var in ("LOCALAPPDATA", "APPDATA", "ProgramFiles", "ProgramFiles(x86)"):
        _base = os.environ.get(_var)
        if _base:
            APP_RESOURCE_CANDIDATES.append(
                os.path.join(_base, "Programs", "DeepSeek Harness", "resources"))
            APP_RESOURCE_CANDIDATES.append(
                os.path.join(_base, "DeepSeek Harness", "resources"))
else:
    APP_RESOURCE_CANDIDATES += [
        os.path.expanduser("~/.local/share/deepseek-harness/resources"),
        "/opt/DeepSeek Harness/resources",
        "/usr/lib/deepseek-harness/resources",
    ]
_dsh_home = os.environ.get("DSH_HOME")
if _dsh_home:
    APP_RESOURCE_CANDIDATES.append(os.path.join(_dsh_home, "resources"))

#: Paths inside an app resource directory. Packaged apps unpack the kit next to
#: `app.asar`; a source checkout keeps it under `dsh/node_modules`.
KIT_NODE_RELATIVE = ("runtime", "primary-runtime", "dependencies", "node", "bin", "node")
KIT_CLI_RELATIVE = ("app.asar.unpacked", "dsh", "node_modules", "@deepseek-ai",
                    "libreoffice-kit", "lib", "cli.js")

#: Fallback searches, so a layout change turns into "found it elsewhere"
#: instead of "LibreOffice is unavailable". Relative to a search root, which is
#: either the app directory or `DshHome`.
KIT_NODE_GLOBS = (
    os.path.join("dependencies", "node", "bin"),
    os.path.join("runtime", "*", "dependencies", "node", "bin"),
    os.path.join("*", "dependencies", "node", "bin"),
)
KIT_CLI_GLOBS = (
    os.path.join("app.asar.unpacked", "dsh", "node_modules", "@deepseek-ai",
                 "libreoffice-kit", "lib", "cli.js"),
    os.path.join("dsh", "node_modules", "@deepseek-ai", "libreoffice-kit", "lib", "cli.js"),
)


# ---------------------------------------------------------------- LibreOffice kit

def _search_roots() -> list[str]:
    """Directories that may hold a DSH application bundle, best guess first."""
    roots = [os.environ.get("DSH_APP_ROOT"), os.environ.get("DSH_RESOURCES")]
    roots += [os.path.dirname(path) for path in _runtimes()]
    roots += APP_RESOURCE_CANDIDATES
    # …/<app>/dsh/assets/wps365-docs/scripts → four levels up is the app root of
    # a source checkout; five levels up covers the packaged layout.
    here = os.path.dirname(os.path.abspath(__file__))
    roots.append(os.path.abspath(os.path.join(here, *([os.pardir] * 4))))
    roots.append(os.path.abspath(os.path.join(here, *([os.pardir] * 5))))
    return [root for root in dict.fromkeys(roots) if root and os.path.isdir(root)]


def _dsh_home() -> str:
    """The DSH home directory, from the environment or the platform default."""
    home = os.environ.get("DSH_HOME")
    return os.path.expanduser(home) if home else os.path.join(os.path.expanduser("~"), ".dsh")


def _runtimes() -> list[str]:
    """Every installed runtime directory under `DshHome/dsh-runtimes`."""
    runtimes = os.path.join(_dsh_home(), "dsh-runtimes")
    try:
        return [os.path.join(runtimes, name) for name in sorted(os.listdir(runtimes))]
    except OSError:
        return []


def find_kit(node_arg: str | None, cli_arg: str | None) -> tuple[str | None, str | None]:
    """定位 LibreOffice kit。优先级：显式参数 → 环境变量 → 已知路径 → PATH。

    应用自带的 kit 在 macOS 上藏在 .app 里，在 Windows 上藏在 resources\\ 下，
    所以两端都按「应用资源目录 + 相对路径」拼；相对路径找不到时再按通配搜一层，
    布局变了也只是"换个地方找到"，而不是报"LibreOffice 不可用"。
    """
    node = node_arg or os.environ.get("DSH_NODE")
    cli = cli_arg or os.environ.get("DSH_LIBREOFFICE_CLI")
    roots = _search_roots()

    if not node:
        for root in roots:
            for leaf in _node_spellings():
                candidate = os.path.join(root, *KIT_NODE_RELATIVE[:-1], leaf)
                if os.path.isfile(candidate):
                    node = candidate
                    break
            if node:
                break
    if not cli:
        for root in roots:
            candidate = os.path.join(root, *KIT_CLI_RELATIVE)
            if os.path.isfile(candidate):
                cli = candidate
                break

    # 通配兜底：上面的相对路径对不上时（打包方式变了），照样能找到 kit。
    # DshHome 下的运行时目录（`dsh-runtimes/<name>/`）既是 node 的家，也放着
    # 应用资源，所以它先于其它 root 被搜。
    if not node or not cli:
        for base in dict.fromkeys([_dsh_home(), *_runtimes(), *roots]):
            if not cli:
                cli = _first_glob(base, KIT_CLI_GLOBS)
            if not node:
                node = _first_glob(base, tuple(
                    os.path.join(pattern, leaf)
                    for pattern in KIT_NODE_GLOBS for leaf in _node_spellings()))
            if node and cli:
                break

    if not node:
        node = shutil.which("node")
    return node, cli


def _node_spellings() -> list[str]:
    """File names `node` may have on this platform (`node`, `node.exe`)."""
    if os.name != "nt":
        return ["node"]
    declared = (os.environ.get("PATHEXT") or ".EXE;.CMD;.BAT").split(";")
    return ["node"] + [f"node{ext.strip().lower()}" for ext in declared if ext.strip()]


def _first_glob(base: str, patterns: tuple[str, ...]) -> str | None:
    """First existing match for any pattern, searched relative to `base`."""
    for pattern in patterns:
        try:
            hits = glob.glob(os.path.join(base, pattern))
        except OSError:
            continue
        if hits:
            return hits[0]
    return None


def recalculate(node: str, cli: str, src: str, dst: str, timeout: int) -> None:
    cmd = [node, cli, "recalculate", "--input", src, "--output", dst]
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout,
                          encoding=CLI_ENCODING, errors="replace")
    if proc.returncode != 0:
        detail = (proc.stderr or proc.stdout or "").strip()[:400]
        sys.exit(f"[xlsx_edit] 重算失败（退出码 {proc.returncode}）：{detail}")
    if not os.path.exists(dst):
        sys.exit("[xlsx_edit] 重算没有产出文件")


# ---------------------------------------------------------------- 编辑

def parse_value(raw: str):
    """把命令行给的字符串还原成合适的单元格类型。"""
    text = raw.strip()
    if not text:
        return None
    if text.startswith("="):                       # 公式（--set 时写成 == 开头）
        return text
    low = text.lower()
    if low in ("true", "false"):
        return low == "true"
    try:
        return int(text)
    except ValueError:
        pass
    try:
        return float(text)
    except ValueError:
        pass
    return raw


def apply_edits(path: str, py_file: str | None, sets: list[str]) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(path)  # data_only=False：公式原文保留
    changed: list[str] = []

    if py_file:
        if not os.path.exists(py_file):
            sys.exit(f"[xlsx_edit] 找不到编辑脚本 {py_file}")
        with open(py_file, encoding="utf-8") as fh:
            code = fh.read()

        def sheet(name: str | None = None):
            return wb[name] if name else wb.active

        env = {"wb": wb, "sheet": sheet, "changes": [], "__name__": "__wps_edit__"}
        try:
            exec(compile(code, py_file, "exec"), env)  # noqa: S102 本地可信脚本
        except Exception as exc:  # noqa: BLE001
            sys.exit(f"[xlsx_edit] 编辑脚本执行失败：{type(exc).__name__}: {exc}")
        changed += [str(x) for x in (env.get("changes") or [])]

    for spec in sets:
        coord, sep, raw = spec.partition("=")
        if not sep:
            sys.exit(f"[xlsx_edit] --set 需要 表!单元格=值 形式，收到：{spec}")
        sheet_name, _, cell = coord.rpartition("!")
        if not sheet_name:
            sys.exit(f"[xlsx_edit] --set 必须写明工作表，收到：{spec}")
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else None
        if ws is None:
            sys.exit(f"[xlsx_edit] 没有工作表 {sheet_name!r}（现有：{wb.sheetnames}）")
        value = parse_value(raw)
        ws[cell] = value
        changed.append(f"{sheet_name}!{cell} = {value}")

    wb.save(path)
    return changed


def count_formulas(path: str) -> tuple[int, int]:
    """返回 (公式数, 其中有缓存值的数量)。"""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    wbv = load_workbook(path, data_only=True)
    total = cached = 0
    for name in wb.sheetnames:
        ws, wsv = wb[name], wbv[name]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    total += 1
                    if wsv[cell.coordinate].value is not None:
                        cached += 1
    return total, cached


# ---------------------------------------------------------------- 读回验证

def literal_expectations(sets: list[str]) -> list[tuple[str, str, str]]:
    """从 `--set` 里挑出能用 sheets 接口核对的字面值：(表名, A1 坐标, 期望值)。

    公式（`=...`）排除在外：sheets 接口只回计算结果、拿不到公式原文，核对不了
    "我写进去的是不是那个公式"。
    """
    final_values: dict[tuple[str, str], str] = {}
    for spec in sets:
        coord, sep, raw = spec.partition("=")
        if not sep:
            continue
        sheet_name, _, cell = coord.rpartition("!")
        if not sheet_name or not cell:
            continue
        final_values[(sheet_name, cell.strip().replace("$", "").upper())] = raw
    # 与 apply_edits 一样，重复赋值以后一次为准，包括最后写入公式的情况。
    return [(sheet, cell, raw) for (sheet, cell), raw in final_values.items()
            if not raw.lstrip().startswith("=")]


def cell_indices(cell: str) -> tuple[int, int]:
    """A1 坐标 → (0-based 行, 0-based 列)，与 sheets 接口的 row_from/col_from 一致。"""
    letters = "".join(ch for ch in cell if ch.isalpha()).upper()
    digits = "".join(ch for ch in cell if ch.isdigit())
    if not letters or not digits:
        raise ValueError(f"无法解析单元格坐标：{cell!r}")
    column = 0
    for ch in letters:
        column = column * 26 + (ord(ch) - ord("A") + 1)
    return int(digits) - 1, column - 1


def same_literal(expected: str, actual) -> bool:
    """宽容比较云端读回值：类型与格式可能与本地写法不同（99 / '99' / 99.0）。"""
    want = parse_value(expected)
    if want is None:
        return actual is None or actual == ""
    if isinstance(want, bool):
        if isinstance(actual, bool):
            return want == actual
        return isinstance(actual, str) and actual.strip().lower() == str(want).lower()
    if isinstance(want, (int, float)):
        if isinstance(actual, bool) or not isinstance(actual, (int, float, str)):
            return False
        try:
            return Decimal(str(want)) == Decimal(str(actual).strip())
        except InvalidOperation:
            return False
    # 引号和非空文本两端的空白都是 parse_value 实际写入的内容。
    return isinstance(actual, str) and want == actual


def verify_adopted(cli: str, file_id: str, expectations: list[tuple[str, str, str]],
                   timeout: int) -> list[str]:
    """把 `--set` 写过的字面值单元格读回来，返回对不上的那些。

    为什么非查不可：**字节覆盖对「在线表格」不生效** —— 版本号与字节都变了，服务端
    内容模型却仍是旧的。2026-10-15 实测：把 B1 从 1 改成 42 后按字节上传，
    `version` 1→2 且云端字节与本机一致，但 sheets 接口读回仍是 1。只信版本号自增
    会把这种情况当成成功（那位同学的原话：「version 增加了，看着像成功」）。
    """
    if not expectations:
        return []
    sheets = file_fetch.api(cli, "get", f"/v7/sheets/{file_id}/worksheets", timeout).get("sheets") or []
    ids = {sheet.get("name"): sheet.get("sheet_id") for sheet in sheets}
    problems: list[str] = []
    for sheet_name, cell, expected in expectations:
        sid = ids.get(sheet_name)
        if sid is None:
            problems.append(f"{sheet_name}!{cell}：云端没有这张工作表（改动未被采纳）")
            continue
        row, col = cell_indices(cell)
        data = file_fetch.api(cli, "get", f"/v7/sheets/{file_id}/worksheets/{sid}/range_data",
                              timeout, query={"row_from": row, "row_to": row,
                                              "col_from": col, "col_to": col})
        actual = None
        for item in data.get("range_data") or []:
            if item.get("row_from") == row and item.get("col_from") == col:
                actual = item.get("original_cell_value")
                if actual in (None, ""):
                    actual = item.get("cell_text")
        # 稀疏响应中缺失的单元格按空值核对，不能把清空操作误报为失败。
        if not same_literal(expected, actual):
            problems.append(f"{sheet_name}!{cell}：期望 {expected!r}，云端读回 {actual!r}")
    return problems


# ---------------------------------------------------------------- 主流程

def main() -> None:
    ap = argparse.ArgumentParser(description="改云端 .xlsx：下载 → 本地编辑 → 重算 → 传回")
    ap.add_argument("--drive-id", required=True, help="文件所在盘 drive_id")
    ap.add_argument("--file-id", required=True, help="要修改的 .xlsx file_id")
    ap.add_argument("--py", help="编辑脚本（以 wb / sheet() 为上下文执行）")
    ap.add_argument("--set", action="append", default=[],
                    metavar="表!单元格=值", help="直接改单元格，可重复；值以 = 开头即公式")
    ap.add_argument("--out", help="工作目录（默认建临时目录；会打印路径）")
    ap.add_argument("--backup-dir", help="上传前把当前版本下载到该目录（建议）")
    ap.add_argument("--no-recalc", action="store_true", help="跳过 LibreOffice 重算")
    ap.add_argument("--dry-run", action="store_true", help="只改本地不传回")
    ap.add_argument("--force", action="store_true", help="忽略冲突检查，强制覆盖")
    ap.add_argument("--node", help="LibreOffice kit 的 node 路径")
    ap.add_argument("--lo-cli", help="LibreOffice kit 的 cli.js 路径")
    ap.add_argument("--cli", default=file_write.DEFAULT_CLI, help="wps365-cli 路径")
    ap.add_argument("--timeout", type=int, default=180, help="单次 HTTP 超时秒数")
    args = ap.parse_args()

    if not args.py and not args.set:
        ap.error("需要 --py 或 --set，否则没有改动")
    if not os.path.exists(args.cli):
        sys.exit(f"[xlsx_edit] 找不到 CLI：{args.cli}")

    work = os.path.abspath(args.out) if args.out else tempfile.mkdtemp(prefix="wps_xlsx_edit_")
    os.makedirs(work, exist_ok=True)

    meta = file_fetch.json_cmd(args.cli, ["drive", "file", "get", args.drive_id, args.file_id],
                               args.timeout, "取文件信息")
    name = meta.get("name") or ""
    if not name.lower().endswith((".xlsx", ".xls")):
        sys.exit(f"[xlsx_edit] {name!r} 不是 .xlsx/.xls —— 本文档请用 file_write.py 整文件传回")
    # 下载那一刻的基线，用于传回前比对（防并发覆盖）
    base_hash = (meta.get("hash") or {}).get("sum")
    base_version = meta.get("version")
    print(f"文件：{name}  version={base_version}  size={meta.get('size')}  "
          f"hash={(base_hash or '')[:12]}…")

    # 1. 下载
    path, nbytes, sha1, declared = file_fetch.download_original(
        args.cli, args.drive_id, args.file_id, work, name, args.timeout)
    print(f"① 已下载 {nbytes} 字节 → {path}")

    # 2. 编辑
    changed = apply_edits(path, args.py, args.set)
    print(f"② 已应用 {len(changed)} 处改动")
    for item in changed[:20]:
        print(f"   · {item}")
    if len(changed) > 20:
        print(f"   … 另有 {len(changed) - 20} 处")

    # 3. 重算（openpyxl 保存会抹掉公式缓存值）
    total, cached = count_formulas(path)
    if total and not args.no_recalc:
        node, lo_cli = find_kit(args.node, args.lo_cli)
        if not node or not lo_cli:
            print(f"③ ⚠️ 未找到 LibreOffice kit，跳过重算 —— {total} 个公式没有缓存值，"
                  "Excel/WPS 打开会重算，但按值读取会拿到空。"
                  "可用 --lo-cli / --node 指定，或设 DSH_LIBREOFFICE_CLI / DSH_NODE。")
        else:
            recalc_path = os.path.join(work, "recalculated" + os.path.splitext(path)[1])
            recalculate(node, lo_cli, path, recalc_path, args.timeout)
            shutil.move(recalc_path, path)
            total, cached = count_formulas(path)
            print(f"③ 已重算：{total} 个公式，{cached} 个有缓存值")
    elif total:
        print(f"③ 跳过重算（--no-recalc）：{total} 个公式，{cached} 个有缓存值")
    else:
        print("③ 无公式，跳过重算")

    print(f"   改后文件：{path}  ({os.path.getsize(path)} 字节)")

    # 4. 传回（传回前先比对：下载之后有没有别人改过）
    if args.dry_run:
        print("④ dry-run：未传回，云端文件未变")
        return
    print("④ 比对云端是否被他人改动…")
    file_write.check_conflict(args.cli, args.drive_id, args.file_id, args.timeout,
                              base_hash, base_version, args.force)
    if args.backup_dir:
        dest = file_write.backup_current(args.cli, args.drive_id, args.file_id,
                                        args.backup_dir, args.timeout, name)
        print(f"   已备份当前版本到 {dest}")
    file_write.upload(args.cli, args.drive_id, meta.get("parent_id") or "0", args.file_id,
                      name, path, None, args.timeout, "delegated", False)

    # 5. 读回验证。字节覆盖对「在线表格」不生效，而它失败得很像成功（版本号照样自增），
    #    所以只信自己读回来的值；读不回一致就明确失败，别让调用方以为改完了。
    expectations = literal_expectations(args.set)
    if not expectations:
        print("⑤ 跳过读回验证：本次改动都是公式或来自 --py —— "
              "sheets 接口读不到公式原文，且 --py 改了什么脚本无从得知")
        return
    print(f"⑤ 读回验证 {len(expectations)} 个用 --set 写的字面值单元格…")
    problems = verify_adopted(args.cli, args.file_id, expectations, args.timeout)
    if problems:
        print("[xlsx_edit] ⚠️ 上传成功，但云端没有采纳这些改动：")
        for item in problems:
            print(f"   · {item}")
        print("   原因：这个文件有服务端内容模型（在线表格）。字节覆盖只对「纯字节文件」生效 ——"
              " 版本号会增加、字节也确实是你的，但在线内容仍是旧的。")
        print("   下一步：改用单元格/公式接口改（见 references/office-files.md 的 .xlsx 接口写入一节），"
              "或让用户在线手动粘贴；**不要**当成已成功。")
        sys.exit(1)
    print(f"   ✅ {len(expectations)} 个单元格读回一致")


if __name__ == "__main__":
    main()
